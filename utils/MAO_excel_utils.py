import pytest
import openpyxl
from playwright.sync_api import Page
from openpyxl import load_workbook


filepath ="data/MAOdatasheet.xlsx"
Maxrow =4

def read_testdata():

    
    workbook = load_workbook(filepath)
    sheet = workbook["Sheet1"]

    data =[]

    for row in range(2, Maxrow+1):
        username =sheet.cell(row,1).value
        password =sheet.cell(row,2).value
        ordervalue =sheet.cell(row,3).value 

        data.append((row,username,password,ordervalue))
    
    workbook.close()
    return data

def write_data(row_num, status, ordertotal, result):
    workbook = load_workbook(filepath)
    sheet = workbook["Sheet1"]

    sheet.cell(row=row_num,column=4).value=status
    sheet.cell(row=row_num,column=5).value=ordertotal
    sheet.cell(row=row_num,column=6).value=result

    workbook.save(filepath)
    workbook.close()
