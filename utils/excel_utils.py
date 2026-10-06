import pytest
import openpyxl
from playwright.sync_api import Page
from openpyxl import load_workbook


filepath ="data/datasheet.xlsx"
Maxrow =4

def read_testdata():

    
    workbook = load_workbook(filepath)
    sheet = workbook["Sheet1"]

    data =[]

    for row in range(2, Maxrow+1):
        search_item =sheet.cell(row,1).value

        data.append((row,search_item))
    
    workbook.close()
    return data

def write_data(row_num, result):
    workbook = load_workbook(filepath)
    sheet = workbook["Sheet1"]

    sheet.cell(row=row_num,column=2).value=result

    workbook.save(filepath)
    workbook.close()
